"""Итоговая таблица каналов из blog_check.csv: чей канал и о чём.

1. Берёт из blog_check.csv только настоящие каналы (status = ok).
2. Для каждого скачивает название, описание и 3 последних поста с t.me/s/.
3. Haiku пачками по 25 определяет: категорию, чей канал, о чём (коротко).
4. Пишет channels_table.xlsx (или channels_table.csv, если нет openpyxl).

Запуск из корня репо (рядом должен лежать blog_check.csv):
    pip install openpyxl     # один раз, если ещё не стоит
    python build_channels_table.py
Займёт ~5–10 минут, Haiku — около $0,3–0,5.
"""
import asyncio
import csv
import json
import re

import httpx
from anthropic import AsyncAnthropic
from bs4 import BeautifulSoup

from src.config import settings

SRC = "blog_check.csv"
MODEL = "claude-haiku-4-5-20251001"
CATEGORIES = [
    "Лента вакансий", "Карьера и поиск работы", "HR и рекрутинг", "Продукт",
    "Маркетинг и PR", "Аналитика и данные", "Дизайн", "Разработка и IT",
    "Менеджмент и лидерство", "Бизнес и предпринимательство", "Другое",
]
UA = {"User-Agent": "Mozilla/5.0"}

SYSTEM = f"""Тебе дают список Telegram-каналов: название, описание и начало последних постов.
Для каждого определи:
- category — ровно одна из: {", ".join(CATEGORIES)};
- owner — чей канал: имя автора и роль/компания, если понятно из описания или постов
  («Анна Петрова, HRD в Ozon»); для медиа или компании — название; если непонятно — «не указано»;
- about — о чём канал, до 15 слов, нейтрально, без оценок.
Ничего не выдумывай: имена и компании — только если они есть в тексте.
Ответ — ТОЛЬКО JSON-массив без пояснений и без ```:
[{{"handle": "...", "category": "...", "owner": "...", "about": "..."}}]"""


async def fetch(client, sem, row):
    handle = row["handle"]
    info = {"handle": handle, "subs": row.get("subs", ""), "period": row.get("period", ""),
            "hiring": row.get("hiring", ""), "hiring_niche": row.get("hiring_niche", ""),
            "title": "", "description": "", "posts": ""}
    async with sem:
        try:
            r = await client.get(f"https://t.me/s/{handle}", headers=UA)
            if r.status_code == 200:
                soup = BeautifulSoup(r.text, "html.parser")
                t = soup.select_one(".tgme_channel_info_header_title")
                d = soup.select_one(".tgme_channel_info_description")
                info["title"] = t.get_text(" ", strip=True) if t else ""
                info["description"] = d.get_text(" ", strip=True)[:400] if d else ""
                texts = [n.get_text(" ", strip=True)[:150] for n in soup.select(".tgme_widget_message_text")]
                info["posts"] = " || ".join(texts[-3:])
        except httpx.HTTPError:
            pass
        await asyncio.sleep(0.4)
    return info


def extract_json(raw: str):
    s, e = raw.find("["), raw.rfind("]")
    if s < 0 or e <= s:
        return []
    try:
        return json.loads(raw[s:e + 1])
    except ValueError:
        return []


async def classify(client, batch):
    user = "\n\n".join(
        f"handle: {c['handle']}\nназвание: {c['title']}\nописание: {c['description']}\nпосты: {c['posts']}"
        for c in batch)
    try:
        resp = await client.messages.create(model=MODEL, max_tokens=4000, system=SYSTEM,
                                            messages=[{"role": "user", "content": user}])
        raw = "".join(b.text for b in resp.content if getattr(b, "type", "") == "text")
        return {it.get("handle", "").lower(): it for it in extract_json(raw) if isinstance(it, dict)}
    except Exception as e:  # noqa: BLE001
        print(f"  ошибка Haiku: {str(e)[:120]}")
        return {}


def subs_num(s: str) -> float:
    s = (s or "").replace(" ", "").upper()
    m = re.match(r"([\d.]+)([KM]?)", s)
    if not m:
        return 0
    return float(m.group(1)) * {"K": 1e3, "M": 1e6, "": 1}[m.group(2)]


async def main():
    with open(SRC, encoding="utf-8") as f:
        rows = [r for r in csv.DictReader(f) if r["status"] == "ok"]
    print(f"Каналов: {len(rows)}. Скачиваю описания…", flush=True)

    sem = asyncio.Semaphore(5)
    async with httpx.AsyncClient(timeout=15) as http:
        channels = await asyncio.gather(*(fetch(http, sem, r) for r in rows))

    print("Размечаю через Haiku…", flush=True)
    ai = AsyncAnthropic(api_key=settings.anthropic_api_key)
    labels = {}
    for i in range(0, len(channels), 25):
        labels.update(await classify(ai, channels[i:i + 25]))
        print(f"  …{min(i + 25, len(channels))}/{len(channels)}", flush=True)

    out = []
    for c in channels:
        lab = labels.get(c["handle"].lower(), {})
        cat = lab.get("category") if lab.get("category") in CATEGORIES else "Другое"
        out.append({
            "Канал": f"https://t.me/{c['handle']}",
            "Название": c["title"],
            "Подписчики": c["subs"],
            "Категория": cat,
            "Чей канал": lab.get("owner", ""),
            "О чём": lab.get("about", ""),
            "Посты о найме": c["hiring"],
            "Из них по нише": c["hiring_niche"],
            "Период последних постов": c["period"],
        })
    out.sort(key=lambda r: (CATEGORIES.index(r["Категория"]), -subs_num(r["Подписчики"])))

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active
        ws.title = "Каналы"
        headers = list(out[0].keys())
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in out:
            ws.append([r[h] for h in headers])
            link = ws.cell(row=ws.max_row, column=1)
            link.hyperlink = r["Канал"]
            link.font = Font(color="0563C1", underline="single")
        for col, width in zip("ABCDEFGHI", [34, 34, 12, 26, 34, 60, 14, 14, 24]):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        wb.save("channels_table.xlsx")
        print("\nГотово: channels_table.xlsx")
    except ImportError:
        with open("channels_table.csv", "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=list(out[0].keys()))
            w.writeheader()
            w.writerows(out)
        print("\nГотово: channels_table.csv (openpyxl не установлен)")

    from collections import Counter
    for cat, n in Counter(r["Категория"] for r in out).most_common():
        print(f"  {cat}: {n}")


if __name__ == "__main__":
    asyncio.run(main())
